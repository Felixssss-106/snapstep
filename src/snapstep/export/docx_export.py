"""Word（docx）导出。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Inches

from ..models import Session


def export_docx(
    session: Session, images: dict[int, Path], out_dir: Path
) -> Path:
    doc = Document()
    doc.add_heading(session.title or "操作教程", level=0)
    if session.intro:
        doc.add_paragraph(session.intro)

    for step in session.steps:
        doc.add_heading(f"{step.index}. {step.title or f'步骤 {step.index}'}", level=2)
        if step.description:
            doc.add_paragraph(step.description)
        typed = step.typed_text()
        if typed:
            p = doc.add_paragraph()
            run = p.add_run(f"输入：{typed}")
            run.italic = True
        if step.has_secret():
            doc.add_paragraph("（本步包含密码输入，内容已隐藏）")
        img = images.get(step.index)
        if img:
            doc.add_picture(str(img), width=Inches(5.9))

    path = out_dir / "guide.docx"
    doc.save(path)
    return path
